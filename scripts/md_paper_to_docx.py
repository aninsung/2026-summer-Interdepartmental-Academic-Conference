"""Convert docs/paper_draft_ko.md into a Word document for pasting/editing."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "paper_draft_ko.md"
OUT = ROOT / "docs" / "paper_draft_ko.docx"

FONT = "Malgun Gothic"
BLUE = "1F4D78"
BLACK = "000000"
GRAY = "555555"
HEADER_FILL = "1F4D78"
ROW_FILL = "F4F7FB"


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), FONT)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def clean_math(text: str) -> str:
    replacements = [
        (r"\\mathrm\{([^}]*)\}", r"\1"),
        (r"\\text\{([^}]*)\}", r"\1"),
        (r"\\frac\{2\|P\\cap G\|\}\{\|P\|\+\|G\|\}", r"2|P ∩ G| / (|P| + |G|)"),
        (r"\\cap", "∩"),
        (r"\\times", "×"),
        (r"\\geq", "≥"),
        (r"\\leq", "≤"),
        (r"\\rightarrow", "→"),
        (r"\\Delta", "Δ"),
        (r"\\, ", " "),
        (r"\\,", " "),
        (r"\^\{-4\}", "⁻⁴"),
        (r"\^\{([^}]*)\}", r"\1"),
        (r"\\left", ""),
        (r"\\right", ""),
        (r"\\,", ""),
    ]
    for pat, repl in replacements:
        text = re.sub(pat, repl, text)
    text = re.sub(r"\\\((.+?)\\\)", r"\1", text)
    text = re.sub(r"\\\[(.+?)\\\]", r"\1", text, flags=re.S)
    text = text.replace("`", "")
    return text


def add_formatted_runs(paragraph, text: str, size=11, color=BLACK, italic=False):
    text = clean_math(text)
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, italic=italic, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, bold=False, italic=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, italic=italic, color=color)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    sizes = {1: 16, 2: 13}
    set_run_font(run, size=sizes.get(level, 13), bold=True, color=BLUE)
    return p


def add_body(doc, text, first_indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.6
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.5)
    add_formatted_runs(p, text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.3
    add_formatted_runs(p, text, size=10, color=GRAY, italic=True)
    return p


def add_list_item(doc, text, numbered=False, index=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    prefix = f"{index}. " if numbered else "• "
    run = p.add_run(prefix)
    set_run_font(run, size=11, bold=numbered)
    add_formatted_runs(p, text)
    return p


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            value = row[j] if j < len(row) else ""
            add_formatted_runs(p, value, size=9, color="FFFFFF" if i == 0 else BLACK)
            if i == 0:
                set_cell_shading(cell, HEADER_FILL)
                for run in p.runs:
                    run.bold = True
            elif i % 2 == 0:
                set_cell_shading(cell, ROW_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def parse_table_row(line: str) -> list[str]:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return [clean_math(c.replace(":---", "").replace("---:", "").replace("---", "")) for c in cells]


def is_table_sep(line: str) -> bool:
    stripped = line.strip().replace("|", "").replace(":", "").replace("-", "").replace(" ", "")
    return line.strip().startswith("|") and stripped == ""


def convert():
    md = SRC.read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    i = 0
    in_formula = False
    formula_lines: list[str] = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if stripped == "\\[" or stripped.startswith("\\["):
            in_formula = True
            formula_lines = [stripped.replace("\\[", "")]
            if "\\]" in stripped:
                in_formula = False
                text = clean_math(" ".join(formula_lines).replace("\\]", ""))
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(8)
                add_formatted_runs(p, text, size=12)
            i += 1
            continue
        if in_formula:
            formula_lines.append(stripped.replace("\\]", ""))
            if "\\]" in stripped:
                in_formula = False
                text = clean_math(" ".join(formula_lines))
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(8)
                add_formatted_runs(p, text, size=12)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            add_formatted_runs(p, stripped[2:], size=18, color=BLUE)
            for run in p.runs:
                run.bold = True
            i += 1
            continue

        if stripped.startswith("## "):
            add_heading(doc, stripped[3:], 1)
            i += 1
            continue

        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 2)
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_table_sep(lines[i]):
                    rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if re.match(r"^\*\*(그림|표)\s+\d+\.\*\*", stripped):
            add_caption(doc, stripped)
            i += 1
            continue

        m_num = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_num:
            add_list_item(doc, m_num.group(2), numbered=True, index=int(m_num.group(1)))
            i += 1
            continue

        if stripped.startswith("- "):
            add_list_item(doc, stripped[2:], numbered=False)
            i += 1
            continue

        add_body(doc, stripped, first_indent=not stripped.startswith("**"))
        i += 1

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    convert()
