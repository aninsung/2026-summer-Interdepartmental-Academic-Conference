import docx
import sys

def extract_text(docx_path, out_path):
    doc = docx.Document(docx_path)
    with open(out_path, 'w', encoding='utf-8') as f:
        for para in doc.paragraphs:
            f.write(para.text + '\n')
        f.write('\n--- TABLES ---\n')
        for table in doc.tables:
            for row in table.rows:
                f.write(' | '.join(cell.text.replace('\n', ' ') for cell in row.cells) + '\n')
            f.write('\n')

if __name__ == '__main__':
    extract_text(sys.argv[1], sys.argv[2])
